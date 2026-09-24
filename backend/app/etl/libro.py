"""Extraer: el libro de Excel convertido en filas tipadas.

Nada de esta etapa mira la base. Solo responde "que dice el libro", con los montos
ya como `Decimal`, las fechas como `date` y cada fila con su numero, para que el
reporte pueda decir "PAGOS, fila 27" y cualquiera la encuentre.
"""

from __future__ import annotations

import hashlib
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path

from openpyxl import load_workbook

#: Version del mapeo columna → campo. Se guarda en `importaciones.mapeo_version`
#: para saber, meses despues, con que reglas se interpreto un archivo.
MAPEO_VERSION = "marfil_v2.1"


class LibroInvalido(ValueError):
    """El archivo no tiene la forma que el ETL sabe leer."""


@dataclass(frozen=True)
class Fila:
    hoja: str
    numero: int
    ref: str
    #: La fila tal cual, para guardarla en `filas_importadas.datos`.
    crudo: dict


@dataclass(frozen=True)
class VentaFila(Fila):
    fecha: date | None
    cliente: str
    producto: str
    precio_usd: Decimal | None
    base: str | None
    costo_usd: Decimal | None
    vendedor: str | None
    plazo_dias: int | None
    tratamiento: str
    observaciones: str


@dataclass(frozen=True)
class PagoFila(Fila):
    fecha: date | None
    venta_ref: str | None
    metodo: str | None
    monto_bs: Decimal | None
    equivalente_usd: Decimal | None
    base: str | None
    referencia: str | None
    concepto: str | None
    tratamiento: str
    observaciones: str


@dataclass(frozen=True)
class EgresoFila(Fila):
    fecha: date | None
    tipo: str
    concepto: str
    monto_usd: Decimal | None
    monto_bs: Decimal | None
    base: str | None
    metodo: str | None
    tratamiento: str
    observaciones: str


@dataclass(frozen=True)
class CostoFila(Fila):
    producto: str
    linea: str | None
    costo_usd: Decimal | None
    es_original: bool


@dataclass
class Libro:
    archivo: str
    hash: str
    ventas: list[VentaFila] = field(default_factory=list)
    pagos: list[PagoFila] = field(default_factory=list)
    egresos: list[EgresoFila] = field(default_factory=list)
    costos: list[CostoFila] = field(default_factory=list)


# ------------------------------------------------------------------ conversores
def _texto(valor: object) -> str | None:
    if valor is None:
        return None
    limpio = str(valor).strip()
    return limpio or None


def _dinero(valor: object) -> Decimal | None:
    """Pasa por `str` y no por `float`: 7466.56 no puede volverse 7466.5599999."""
    if valor is None or (isinstance(valor, str) and not valor.strip()):
        return None
    try:
        monto = Decimal(str(valor).replace(",", "."))
    except InvalidOperation:
        return None
    return monto.quantize(Decimal("0.01"))


def _fecha(valor: object) -> date | None:
    if isinstance(valor, datetime):
        return valor.date()
    if isinstance(valor, date):
        return valor
    return None


def _entero(valor: object) -> int | None:
    try:
        return int(valor) if valor is not None else None
    except (TypeError, ValueError):
        return None


def _crudo(valor: object) -> object:
    if isinstance(valor, (datetime, date)):
        return valor.isoformat()[:10]
    if isinstance(valor, Decimal):
        return str(valor)
    return valor


# ------------------------------------------------------------------- lectura
def _tabla(libro, hoja: str, primera_columna: str) -> list[tuple[int, dict]]:
    """Filas de una hoja como `{cabecera: valor}`, ubicando la cabecera por nombre.

    El libro tiene un titulo y una linea de instrucciones antes de la cabecera, y
    eso puede cambiar entre versiones; buscar la fila cuya primera celda es
    `primera_columna` es mas robusto que asumir que la cabecera esta en la fila 3.
    """
    if hoja not in libro.sheetnames:
        raise LibroInvalido(f"Falta la hoja {hoja!r}.")
    filas = list(libro[hoja].iter_rows(values_only=True))
    try:
        indice = next(
            i for i, f in enumerate(filas) if f and _texto(f[0]) == primera_columna
        )
    except StopIteration as exc:
        raise LibroInvalido(
            f"La hoja {hoja!r} no tiene una cabecera que empiece por {primera_columna!r}."
        ) from exc
    cabecera = [_texto(c) for c in filas[indice]]
    salida = []
    for desplazamiento, valores in enumerate(filas[indice + 1 :], start=indice + 2):
        registro = {c: v for c, v in zip(cabecera, valores, strict=False) if c}
        salida.append((desplazamiento, registro))
    return salida


def _obs(registro: dict) -> str:
    partes = [_texto(registro.get("Observaciones")), _texto(registro.get("Origen"))]
    return " · ".join(p for p in partes if p)


def leer(ruta: str | Path) -> Libro:
    ruta = Path(ruta)
    contenido = ruta.read_bytes()
    libro = load_workbook(ruta, data_only=True, read_only=True)
    resultado = Libro(archivo=ruta.name, hash=hashlib.sha256(contenido).hexdigest())

    for numero, r in _tabla(libro, "VENTAS", "ID venta"):
        ref, cliente = _texto(r.get("ID venta")), _texto(r.get("Cliente"))
        # Fila de plantilla: el ID ya viene generado pero no hay venta.
        if not ref or not cliente:
            continue
        resultado.ventas.append(
            VentaFila(
                hoja="VENTAS",
                numero=numero,
                ref=ref,
                crudo={k: _crudo(v) for k, v in r.items()},
                fecha=_fecha(r.get("Fecha venta")),
                cliente=cliente,
                producto=_texto(r.get("Producto")) or "",
                precio_usd=_dinero(r.get("Precio pactado $")),
                base=_texto(r.get("Base del precio")),
                costo_usd=_dinero(r.get("Costo histórico $")),
                vendedor=_texto(r.get("Vendedor")),
                plazo_dias=_entero(r.get("Plazo días")),
                tratamiento=_texto(r.get("Tratamiento")) or "Incluir",
                observaciones=_obs(r),
            )
        )

    for numero, r in _tabla(libro, "PAGOS", "ID pago"):
        ref, fecha = _texto(r.get("ID pago")), _fecha(r.get("Fecha cobro"))
        if not ref or fecha is None:
            continue
        referencia = _texto(r.get("Referencia"))
        resultado.pagos.append(
            PagoFila(
                hoja="PAGOS",
                numero=numero,
                ref=ref,
                crudo={k: _crudo(v) for k, v in r.items()},
                fecha=fecha,
                venta_ref=_texto(r.get("ID venta")),
                metodo=_texto(r.get("Método")),
                monto_bs=_dinero(r.get("Recibido Bs")),
                equivalente_usd=_dinero(r.get("Equivalente acordado $")),
                base=_texto(r.get("Base")),
                referencia=referencia,
                concepto=_texto(r.get("Concepto")),
                tratamiento=_texto(r.get("Tratamiento")) or "Incluir",
                observaciones=_obs(r),
            )
        )

    for numero, r in _tabla(libro, "GASTOS Y COMPRAS", "ID egreso"):
        ref, fecha = _texto(r.get("ID egreso")), _fecha(r.get("Fecha"))
        if not ref or fecha is None:
            continue
        resultado.egresos.append(
            EgresoFila(
                hoja="GASTOS Y COMPRAS",
                numero=numero,
                ref=ref,
                crudo={k: _crudo(v) for k, v in r.items()},
                fecha=fecha,
                tipo=_texto(r.get("Tipo")) or "Gasto",
                concepto=_texto(r.get("Concepto / lote")) or "",
                monto_usd=_dinero(r.get("Importe $ ref.")),
                monto_bs=_dinero(r.get("Importe Bs")),
                base=_texto(r.get("Base")),
                metodo=_texto(r.get("Método")),
                tratamiento=_texto(r.get("Tratamiento")) or "Incluir",
                observaciones=_obs(r),
            )
        )

    for hoja, columna_producto, columna_costo, es_original in (
        ("PRECIOS PERF TOP QUALITY", "Producto", "Costo $", False),
        ("PRECIOS PERF ORIGINALES", "Producto original", "Precio original $", True),
    ):
        if hoja not in libro.sheetnames:
            continue
        for numero, r in _tabla(libro, hoja, "Código"):
            ref, producto = _texto(r.get("Código")), _texto(r.get(columna_producto))
            if not ref or not producto:
                continue
            resultado.costos.append(
                CostoFila(
                    hoja=hoja,
                    numero=numero,
                    ref=ref,
                    crudo={k: _crudo(v) for k, v in r.items()},
                    producto=producto,
                    linea=_texto(r.get("Línea")),
                    costo_usd=_dinero(r.get(columna_costo)),
                    es_original=es_original,
                )
            )

    libro.close()
    return resultado
